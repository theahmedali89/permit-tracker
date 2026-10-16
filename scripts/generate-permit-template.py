"""Generate the free permit-tracking spreadsheet template.

Output: public/templates/permit-tracking-template.xlsx
Sheet "Permits" with a banner row, a how-to-use note, column headers,
sample rows, and real conditional formatting on the Expiry Date column:
  green  = more than 60 days left
  yellow = 30-60 days left
  red    = under 30 days left or expired
"""
from __future__ import annotations

from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUT_PATH = Path(__file__).resolve().parent.parent / "public" / "templates" / "permit-tracking-template.xlsx"

HEADERS = ["Permit #", "Type", "Jurisdiction", "Issue Date", "Expiry Date", "Status", "Notes"]

SAMPLE_ROWS = [
    ["BLD-2026-0142", "Building", "City of Austin", date.today() - timedelta(days=40), date.today() + timedelta(days=140), "Active", "Sample: delete or overwrite"],
    ["ELE-2026-0087", "Electrical", "City of Denver", date.today() - timedelta(days=200), date.today() + timedelta(days=45), "Active", "Sample: delete or overwrite"],
    ["MEC-2025-0301", "Mechanical", "Harris County", date.today() - timedelta(days=340), date.today() - timedelta(days=10), "Expired", "Sample: delete or overwrite"],
]


def build() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Permits"

    # Row 1: banner
    ws.merge_cells("A1:G1")
    banner = ws["A1"]
    banner.value = "Tired of manual entry? Try Permit Warden free \u2014 permitwarden.com"
    banner.font = Font(bold=True, size=13, color="FFFFFF")
    banner.fill = PatternFill("solid", fgColor="1D4ED8")
    banner.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    # Row 2: how-to-use note
    ws.merge_cells("A2:G2")
    note = ws["A2"]
    note.value = (
        "How to use: fill one row per permit/license/inspection. "
        "The Expiry Date column is color-coded automatically \u2014 green = more than 60 days left, "
        "yellow = 30\u201360 days left, red = under 30 days left or expired. "
        "Sort or filter by Expiry Date to see what's due next. Delete the sample rows when you're ready."
    )
    note.font = Font(size=10, italic=True)
    note.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 44

    # Row 3: headers
    for col, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=3, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="334155")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Sample data rows
    for r, row in enumerate(SAMPLE_ROWS, start=4):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            if c in (4, 5):  # date columns
                cell.number_format = "MM/DD/YYYY"

    # Conditional formatting on Expiry Date (column E), data rows and beyond
    expiry_col = get_column_letter(5)
    target = f"{expiry_col}4:{expiry_col}2000"
    green = PatternFill("solid", fgColor="C6EFCE")
    yellow = PatternFill("solid", fgColor="FFEB9C")
    red = PatternFill("solid", fgColor="FFC7CE")
    ws.conditional_formatting.add(
        target,
        FormulaRule(formula=[f'AND({expiry_col}4<>"",{expiry_col}4<TODAY()+30)'], fill=red),
    )
    ws.conditional_formatting.add(
        target,
        FormulaRule(
            formula=[f'AND({expiry_col}4<>"",{expiry_col}4>=TODAY()+30,{expiry_col}4<=TODAY()+60)'],
            fill=yellow,
        ),
    )
    ws.conditional_formatting.add(
        target,
        FormulaRule(formula=[f'AND({expiry_col}4<>"",{expiry_col}4>TODAY()+60)'], fill=green),
    )

    widths = [16, 16, 20, 13, 13, 12, 40]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:G3"

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_PATH)
    print(f"Wrote {OUT_PATH} ({OUT_PATH.stat().st_size} bytes)")


if __name__ == "__main__":
    build()
